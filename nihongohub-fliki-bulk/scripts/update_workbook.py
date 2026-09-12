"""Regenerate the distributable workbook from the validated CSV."""

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from models import FIELDS, load_quiz, load_rows


def main():
    source = ROOT / "quiz_10.csv"
    load_quiz(source)
    book = Workbook()
    sheet = book.active
    sheet.title = "Quiz Data"
    sheet.append(FIELDS)
    numeric = {
        "question_number",
        "question_page",
        "answer_page",
        "question_start",
        "answer_start",
    }
    for row in load_rows(source):
        sheet.append([int(row[k]) if k in numeric else row[k] for k in FIELDS])
    for cell in sheet[1]:
        cell.font = Font(name="Calibri", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="C64D18")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[1].height = 34
    for row in sheet.iter_rows(min_row=2):
        sheet.row_dimensions[row[0].row].height = 44
        for cell in row:
            cell.font = Font(name="Calibri", size=11)
            cell.alignment = Alignment(wrap_text=True, vertical="center")
    for i, key in enumerate(FIELDS, 1):
        sheet.column_dimensions[get_column_letter(i)].width = (
            42
            if "voice_text" in key or key == "question_text"
            else 20
            if "option_" in key or "correct_" in key
            else 17
        )
    sheet.freeze_panes = "E2"
    table = Table(displayName="NihongoHubQuiz", ref=f"A1:S{sheet.max_row}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    sheet.add_table(table)
    sheet.sheet_view.showGridLines = False
    destination = ROOT / "NihongoHub_Fliki_Bulk_10.xlsx"
    book.save(destination)
    assert load_quiz(destination) == load_quiz(source)
    print(destination)


if __name__ == "__main__":
    main()
